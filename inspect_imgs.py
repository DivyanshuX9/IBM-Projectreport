import docx

doc = docx.Document('Assigment2 - Divyanshu Sharma.docx')
img_idx = 0
for i, p in enumerate(doc.paragraphs):
    has_img = False
    for r in p.runs:
        if 'graphic' in r._element.xml:
            has_img = True
    if has_img:
        img_idx += 1
        prev_text = doc.paragraphs[i-1].text.strip() if i > 0 else ''
        next_text = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ''
        print(f'Image {img_idx} (near P{i}): prev="{prev_text[:60]}", next="{next_text[:60]}"')

