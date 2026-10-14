import docx

doc = docx.Document('Assigment2 - Divyanshu Sharma.docx')
img_num = 0
for i, p in enumerate(doc.paragraphs):
    has_img = any('graphic' in r._element.xml for r in p.runs)
    if has_img:
        img_num += 1
        # walk backwards to find heading or strong text
        heading = ""
        for j in range(i, max(-1, i-15), -1):
            t = doc.paragraphs[j].text.strip()
            if t.startswith("Part") or t.startswith("Phase") or t.startswith("Mini Project") or t.startswith("Figure") or t.startswith("Fig") or "Prompt" in t or "Colab" in t or t.startswith("1.") or t.startswith("2.") or t.startswith("3.") or t.startswith("4.") or t.startswith("5.") or t.startswith("6.") or t.startswith("7.") or t.startswith("8."):
                heading = t
                break
        print(f"Img {img_num:2d} | Heading: {heading[:70]}")

