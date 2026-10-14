import docx

doc_template = docx.Document('Summer Internship Report Template-BDS.docx')
print("=== TEMPLATE TABLES ===")
for t_idx, t in enumerate(doc_template.tables):
    txt = ' '.join(c.text for row in t.rows for c in row.cells)
    if 'CHAPTER' in txt.upper() or 'CONTENTS' in txt.upper() or 'TABLE OF' in txt.upper():
        print(f"Table {t_idx} (rows: {len(t.rows)}, cols: {len(t.columns)}):")
        for r_idx, row in enumerate(t.rows):
            cell_texts = [c.text.strip().replace('\n', ' ') for c in row.cells]
            print(f"  Row {r_idx:2d}: {cell_texts}")

print("\n=== TEMPLATE PARAGRAPHS WITH CONTENTS ===")
for p_idx, p in enumerate(doc_template.paragraphs):
    if 'CONTENTS' in p.text.upper():
        print(f"P{p_idx}: {p.text}")
        for j in range(p_idx, min(len(doc_template.paragraphs), p_idx+25)):
            print(f"  P{j}: {repr(doc_template.paragraphs[j].text)}")

