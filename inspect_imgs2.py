import docx

doc = docx.Document('Assigment2 - Divyanshu Sharma.docx')
img_idx = 0
for i, p in enumerate(doc.paragraphs):
    has_img = any('graphic' in r._element.xml for r in p.runs)
    if has_img:
        img_idx += 1
        # find closest non-empty text above
        above_text = ""
        for j in range(i-1, max(-1, i-6), -1):
            if doc.paragraphs[j].text.strip():
                above_text = doc.paragraphs[j].text.strip()
                break
        print(f"Image {img_idx:2d} (P{i:3d}): {above_text[:90]}")

