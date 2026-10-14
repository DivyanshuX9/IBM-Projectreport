import docx

doc = docx.Document('Summer Internship Report Template-BDS.docx')
t0 = doc.tables[0]
print("=== TEMPLATE TABLE 0 XML ===")
print("Table XML (first 1000 chars):")
print(t0._element.xml[:1000])

print("\n=== CELL WIDTHS & PROPERTIES ===")
for r_idx in [0, 2, 13]:
    row = t0.rows[r_idx]
    for c_idx, cell in enumerate(row.cells):
        print(f"Row {r_idx}, Col {c_idx}: width={cell.width}, text={repr(cell.text.strip())}")
        for p in cell.paragraphs:
            pPr = p._p.pPr
            print(f"   pStyle={p.style.name if p.style else None}, align={p.alignment}, runs={[r.text for r in p.runs]}")


